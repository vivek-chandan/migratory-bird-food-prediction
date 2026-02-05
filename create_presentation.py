#!/usr/bin/env python3
"""
Script to create a PowerPoint presentation for the Migratory Bird Food Prediction project
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_title_slide(prs):
    """Create the title slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title slide layout
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Migratory Bird Stopover Quality Index"
    
    subtitle_frame = subtitle.text_frame
    subtitle_frame.clear()
    
    p1 = subtitle_frame.paragraphs[0]
    p1.text = "Project Update Presentation"
    p1.font.size = Pt(24)
    p1.font.bold = True
    
    # Add team information
    p2 = subtitle_frame.add_paragraph()
    p2.text = "\n\nPresented By:"
    p2.font.size = Pt(18)
    p2.font.bold = True
    
    p3 = subtitle_frame.add_paragraph()
    p3.text = "Vivek Kumar Chandan (2024201032)"
    p3.font.size = Pt(16)
    
    p4 = subtitle_frame.add_paragraph()
    p4.text = "Abhishek Gupta (2024201050)"
    p4.font.size = Pt(16)
    
    p5 = subtitle_frame.add_paragraph()
    p5.text = "\n\nSupervisor: Dr. Rama Chandra Pillutla"
    p5.font.size = Pt(16)
    p5.font.italic = True
    
    p6 = subtitle_frame.add_paragraph()
    p6.text = "IIIT Hyderabad"
    p6.font.size = Pt(16)
    p6.font.italic = True

def create_overview_slide(prs):
    """Create project overview slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    title = slide.shapes.title
    title.text = "Project Overview"
    
    # Add content
    left = Inches(1)
    top = Inches(1.5)
    width = Inches(8)
    height = Inches(4.5)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Objective:"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    p = tf.add_paragraph()
    p.text = "Develop a Migratory Bird Stopover Quality Index to quantify food and habitat availability at stopover sites during peak migration months."
    p.font.size = Pt(16)
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "\nKey Goals:"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    goals = [
        "Analyze eBird migratory species data with NDVI vegetation indices",
        "Identify critical stopover locations for migratory birds",
        "Quantify habitat quality based on vegetation greenness",
        "Detect declining stopover zones",
        "Focus on understudied areas in global migration maps"
    ]
    
    for goal in goals:
        p = tf.add_paragraph()
        p.text = goal
        p.font.size = Pt(14)
        p.level = 1

def create_dataset_collection_slide(prs):
    """Create dataset collection status slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Dataset Collection Status"
    
    left = Inches(1)
    top = Inches(1.5)
    width = Inches(8)
    height = Inches(5)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Primary Datasets Identified:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    datasets = [
        ("eBird Basic Dataset (EBD)", "Global bird observation data from Cornell Lab", "✓ Access Obtained"),
        ("MODIS NDVI (MOD13Q1)", "250m vegetation indices from NASA", "✓ Access Configured"),
        ("Important Bird Areas (IBA)", "Critical stopover sites from BirdLife International", "✓ Documented"),
        ("Landsat/Sentinel-2 NDVI", "High-resolution vegetation data (30m/10m)", "✓ Available"),
        ("ESA Land Cover Data", "Habitat classification for stopover sites", "✓ Identified"),
        ("WorldClim", "Climate data for environmental analysis", "✓ Ready")
    ]
    
    for name, desc, status in datasets:
        p = tf.add_paragraph()
        p.text = f"{name}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.level = 1
        
        p = tf.add_paragraph()
        p.text = f"{desc} - {status}"
        p.font.size = Pt(12)
        p.level = 2

def create_attributes_slide(prs):
    """Create key attributes slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Key Data Attributes"
    
    left = Inches(0.8)
    top = Inches(1.5)
    width = Inches(8.5)
    height = Inches(5)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    # eBird Data Attributes
    p = tf.paragraphs[0]
    p.text = "eBird Data Attributes:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(204, 0, 0)
    
    ebird_attrs = [
        "Species Name (Common & Scientific)",
        "Observation Date & Time (Migration timing)",
        "Geographic Coordinates (Latitude, Longitude)",
        "Bird Count (Abundance data)",
        "Observer Effort (Protocol type, duration)"
    ]
    
    for attr in ebird_attrs:
        p = tf.add_paragraph()
        p.text = attr
        p.font.size = Pt(12)
        p.level = 1
    
    # NDVI Data Attributes
    p = tf.add_paragraph()
    p.text = "\nNDVI Data Attributes:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 153, 0)
    
    ndvi_attrs = [
        "Vegetation Index Values (NDVI, EVI)",
        "Spatial Resolution (10m - 250m)",
        "Temporal Coverage (16-day composites)",
        "Quality Assurance Flags",
        "Time Series Data (2000-present)"
    ]
    
    for attr in ndvi_attrs:
        p = tf.add_paragraph()
        p.text = attr
        p.font.size = Pt(12)
        p.level = 1
    
    # Habitat Data
    p = tf.add_paragraph()
    p.text = "\nHabitat & Environmental Attributes:"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 102, 204)
    
    habitat_attrs = [
        "Land Cover Classification (22 classes)",
        "Protected Area Status",
        "Climate Variables (Temperature, Precipitation)"
    ]
    
    for attr in habitat_attrs:
        p = tf.add_paragraph()
        p.text = attr
        p.font.size = Pt(12)
        p.level = 1

def create_approach_slide(prs):
    """Create methodology/approach slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Our Approach & Methodology"
    
    left = Inches(1)
    top = Inches(1.5)
    width = Inches(8)
    height = Inches(5)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Data Processing Workflow:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    steps = [
        ("1. Data Acquisition", "Download and filter eBird observations for migratory species during peak migration months (March-May, August-October)"),
        ("2. Stopover Site Identification", "Apply spatial clustering to identify key stopover locations from bird observation density"),
        ("3. NDVI Extraction", "Extract vegetation indices at stopover sites during migration periods using Google Earth Engine"),
        ("4. Quality Index Calculation", "Combine NDVI (food availability) with bird abundance to compute stopover quality metrics"),
        ("5. Temporal Analysis", "Analyze multi-year trends to identify declining stopover zones"),
        ("6. Validation & Reporting", "Cross-reference with Important Bird Areas database and conservation literature")
    ]
    
    for step, desc in steps:
        p = tf.add_paragraph()
        p.text = step
        p.font.size = Pt(14)
        p.font.bold = True
        p.level = 1
        
        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.level = 2

def create_next_steps_slide(prs):
    """Create next steps slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Next Steps & Timeline"
    
    left = Inches(1)
    top = Inches(1.5)
    width = Inches(8)
    height = Inches(5)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Immediate Next Steps:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    next_steps = [
        ("Week 1-2: Data Download & Setup", [
            "Complete eBird data request and download filtered datasets",
            "Set up Google Earth Engine authentication",
            "Download MODIS NDVI data for pilot regions"
        ]),
        ("Week 3-4: Data Processing Pipeline", [
            "Develop scripts for eBird data filtering (migration periods)",
            "Implement stopover site clustering algorithm",
            "Create NDVI extraction workflow"
        ]),
        ("Week 5-6: Analysis & Index Development", [
            "Calculate stopover quality indices",
            "Perform statistical analysis on trends",
            "Identify declining zones"
        ]),
        ("Week 7-8: Validation & Documentation", [
            "Validate results with known Important Bird Areas",
            "Document findings and create visualizations",
            "Prepare final report and presentation"
        ])
    ]
    
    for phase, tasks in next_steps:
        p = tf.add_paragraph()
        p.text = phase
        p.font.size = Pt(13)
        p.font.bold = True
        p.level = 1
        
        for task in tasks:
            p = tf.add_paragraph()
            p.text = task
            p.font.size = Pt(10)
            p.level = 2

def create_summary_slide(prs):
    """Create summary/conclusion slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Summary & Expected Outcomes"
    
    left = Inches(1)
    top = Inches(1.5)
    width = Inches(8)
    height = Inches(5)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Current Status:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 102, 0)
    
    status_items = [
        "✓ Comprehensive dataset sources identified and documented",
        "✓ Data access mechanisms configured (eBird, NASA, Google Earth Engine)",
        "✓ Analysis workflow designed and ready for implementation",
        "✓ Development environment and tools set up"
    ]
    
    for item in status_items:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)
        p.level = 1
    
    p = tf.add_paragraph()
    p.text = "\nExpected Deliverables:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)
    
    deliverables = [
        "Stopover Quality Index for migratory bird species",
        "Interactive maps showing critical stopover locations",
        "Temporal analysis identifying declining zones",
        "Conservation recommendations for understudied regions",
        "Complete documentation and reproducible code"
    ]
    
    for item in deliverables:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.level = 1
    
    # Add thank you note
    p = tf.add_paragraph()
    p.text = "\n\nThank You!"
    p.font.size = Pt(24)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

def main():
    """Main function to create the presentation"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    print("Creating presentation slides...")
    
    create_title_slide(prs)
    print("✓ Slide 1: Title slide created")
    
    create_overview_slide(prs)
    print("✓ Slide 2: Project overview created")
    
    create_dataset_collection_slide(prs)
    print("✓ Slide 3: Dataset collection status created")
    
    create_attributes_slide(prs)
    print("✓ Slide 4: Key attributes created")
    
    create_approach_slide(prs)
    print("✓ Slide 5: Approach and methodology created")
    
    create_next_steps_slide(prs)
    print("✓ Slide 6: Next steps and timeline created")
    
    create_summary_slide(prs)
    print("✓ Slide 7: Summary and outcomes created")
    
    # Save the presentation
    output_file = "Migratory_Bird_Project_Update.pptx"
    prs.save(output_file)
    print(f"\n✓ Presentation saved as: {output_file}")
    print(f"Total slides created: {len(prs.slides)}")

if __name__ == "__main__":
    main()
